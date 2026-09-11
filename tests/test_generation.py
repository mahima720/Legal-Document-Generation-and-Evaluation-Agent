import pytest
from pathlib import Path
from docx import Document
from src.config import FILE_CASE_INFO, OUTPUTS_DIR
from src.ingestion.case_parser import parse_case_information
from src.mapping.mapper import map_case_to_content_plan
from src.generation.drafter import AffidavitDrafter
from src.generation.docx_renderer import render_affidavit_to_docx
from src.schemas.document import AffidavitDocument


@pytest.fixture
def content_plan():
    case_info = parse_case_information(FILE_CASE_INFO)
    return map_case_to_content_plan(case_info)


@pytest.fixture
def drafter():
    return AffidavitDrafter()


def test_mock_drafting_produces_valid_affidavit_document(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    assert isinstance(doc, AffidavitDocument)
    assert doc.court_heading == "IN THE HIGH COURT OF JUDICATURE AT BOMBAY"
    assert doc.jurisdiction == "ORDINARY ORIGINAL CIVIL JURISDICTION"
    assert "1847 OF 2026" in doc.case_number_line
    # Invariant: No sample case leakage
    assert "3147" not in doc.case_number_line


def test_paragraph_numbering_sequential_and_count_is_7(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    assert len(doc.body_paragraphs) == 7
    for idx, p in enumerate(doc.body_paragraphs, start=1):
        assert p.number == idx
        assert len(p.text) > 20


def test_verification_range_is_paragraphs_1_to_7(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    assert doc.verification.paragraph_range_start == 1
    assert doc.verification.paragraph_range_end == 7
    assert "paragraphs 1 to 7" in doc.verification.verification_text.lower()
    # Invariant: Must not copy sample affidavit range
    assert "paragraphs 1 to 5" not in doc.verification.verification_text.lower()


def test_respondent_and_deponent_info_correct(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    assert "Arvind Rajan" in doc.deponent_clause
    assert "Deputy Metropolitan Commissioner" in doc.deponent_clause
    assert "RESPONDENT NO. 2" in doc.affidavit_title.upper() or "RESPONDENT NO.2" in doc.affidavit_title.upper()
    assert "Sunrise Housing Private Limited" in doc.cause_title_petitioner

    # Invariant: Sample parties must never leak
    assert "Arjun Mehta" not in doc.cause_title_petitioner
    assert "Rohan Deshpande" not in doc.deponent_clause


def test_exhibit_reference_and_prayer_present(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    exhibit_found = any("EXHIBIT-‘A’" in p.text or "EXHIBIT-'A'" in p.text for p in doc.body_paragraphs)
    assert exhibit_found is True

    assert len(doc.prayer_clauses) == 3
    assert doc.prayer_clauses[0].letter == "(a)"
    assert doc.prayer_clauses[1].letter == "(b)"
    assert doc.prayer_clauses[2].letter == "(c)"


def test_docx_rendering_and_content_verification(drafter, content_plan):
    doc = drafter.draft(content_plan, mode="mock")
    docx_path = OUTPUTS_DIR / "generated_affidavit.docx"
    result_path = render_affidavit_to_docx(doc, docx_path)

    assert result_path.exists()
    assert result_path.stat().st_size > 5000

    # Read back and inspect DOCX content
    docx_file = Document(result_path)
    para_texts = [p.text for p in docx_file.paragraphs]
    table_texts = [c.text for t in docx_file.tables for row in t.rows for c in row.cells]
    combined_text = " ".join(para_texts + table_texts)

    # Required major sections
    assert "IN THE HIGH COURT OF JUDICATURE AT BOMBAY" in combined_text
    assert "ORDINARY ORIGINAL CIVIL JURISDICTION" in combined_text
    assert "1847 OF 2026" in combined_text
    assert "Sunrise Housing Private Limited" in combined_text
    assert "Mumbai Metropolitan Region Development Authority" in combined_text
    assert "Arvind Rajan" in combined_text
    assert "Deputy Metropolitan Commissioner" in combined_text
    assert "PRAYER" in combined_text
    assert "VERIFICATION" in combined_text
    assert "RAJAN & ASSOCIATES" in combined_text

    # Critical paragraph rule in verification
    assert "paragraphs 1 to 7" in combined_text
    assert "paragraphs 1 to 5" not in combined_text

    # No sample leakage
    assert "3147 OF 2026" not in combined_text
    assert "Rohan Deshpande" not in combined_text
    assert "Arjun Mehta" not in combined_text
